"""Architecture contracts: pure imports, isolated composition and package boundaries."""
import ast
import asyncio
from dataclasses import FrozenInstanceError
import json
from pathlib import Path
import subprocess
import sys

import pytest

from xmind_mcp.models import Document, Topic, ReadLimits, normalize_topics, walk_topics
from xmind_mcp.parsers import BaseParser
from xmind_mcp.parsers.registry import ParserRegistry
from xmind_mcp.service import XMindService
from xmind_mcp.settings import Settings
from xmind_mcp.xmind_io.codec import document_to_sheet
from xmind_mcp.xmind_io.reader import XMindReader
from xmind_mcp.xmind_io.writer import create_xmind_file

ROOT = Path(__file__).resolve().parents[1]


def test_package_imports_have_no_process_side_effects(tmp_path):
    code = f'''
import sys, os, logging
sys.path.insert(0, {str(ROOT)!r})
original_path = list(sys.path)
original_cwd = os.getcwd()
original_handlers = list(logging.getLogger().handlers)
original_level = logging.getLogger().level
import xmind_mcp
import xmind_mcp.server
import xmind_mcp.service
assert os.getcwd() == original_cwd
assert sys.path == original_path
assert logging.getLogger().handlers == original_handlers
assert logging.getLogger().level == original_level
assert 'mcp' not in sys.modules
'''
    result = subprocess.run([sys.executable, '-I', '-c', code], cwd=tmp_path,
                            capture_output=True, text=True, timeout=20)
    assert result.returncode == 0, result.stderr
    assert result.stdout == result.stderr == ''


def test_package_has_no_root_module_dependencies():
    legacy = {'xmind_core_engine', 'xmind_mcp_server', 'universal_xmind_converter',
              'xmind_writer', 'validate_xmind_structure', 'xmind_types', 'xmind_ai_extensions'}
    assert all(not (ROOT / f'{name}.py').exists() for name in legacy)
    for file in (ROOT / 'xmind_mcp').rglob('*.py'):
        for node in ast.walk(ast.parse(file.read_text())):
            if isinstance(node, ast.Import):
                assert not any(name.name.split('.')[0] in legacy for name in node.names), file
            elif isinstance(node, ast.ImportFrom) and not node.level:
                assert (node.module or '').split('.')[0] not in legacy, file


@pytest.mark.parametrize('name', ['test_txt.txt', 'test_markdown.md', 'test_html.html', 'test_word.docx', 'test_excel.xlsx'])
def test_all_input_adapters_return_canonical_documents(name):
    document = ParserRegistry().get_parser(ROOT / 'examples' / name).parse()
    assert isinstance(document, Document)
    assert isinstance(document.root, Topic)
    assert list(walk_topics(document.root))
    assert all(isinstance(topic, Topic) for topic, _ in walk_topics(document.root))


def test_environment_is_captured_once_and_settings_are_immutable(tmp_path, monkeypatch):
    env = {'XMIND_MCP_BASE_DIR': str(tmp_path), 'XMIND_MCP_DEFAULT_OUTPUT_DIR': 'from-env',
           'XMIND_MAX_NODE_COUNT': '12'}
    settings = Settings.from_env('from-cli', env)
    assert settings.default_output_dir == tmp_path / 'from-cli'
    env['XMIND_MAX_NODE_COUNT'] = '90'
    assert settings.read_limits.max_nodes == 12
    monkeypatch.chdir(tmp_path.parent)
    assert settings.input_path('relative.md') == tmp_path / 'relative.md'
    with pytest.raises(FrozenInstanceError):
        settings.default_output_dir = tmp_path / 'changed'


@pytest.mark.parametrize('values', [(0, 1, 1), (1, 0, 1), (1, 1, -1)])
def test_invalid_read_limits_fail_at_configuration_boundary(values):
    with pytest.raises(ValueError):
        ReadLimits(*values)


def test_custom_parser_is_local_to_one_application(tmp_path):
    class CustomParser(BaseParser):
        def parse(self):
            return Document(Topic('custom document'))

    first = XMindService(Settings(tmp_path, tmp_path / 'first'))
    second = XMindService(Settings(tmp_path, tmp_path / 'second'))
    first.parser_factory.register('.custom', CustomParser)
    source = tmp_path / 'input.custom'
    source.write_text('custom source')
    assert first.convert_to_xmind(str(source))['status'] == 'success'
    assert second.convert_to_xmind(str(source))['status'] == 'error'
    assert not (tmp_path / 'second/input.xmind').exists()


@pytest.mark.parametrize('mode', ['auto', 'legacy'])
def test_two_servers_keep_configuration_and_results_isolated(tmp_path, mode):
    from mcp import Client
    from xmind_mcp.server import create_server
    first = tmp_path / 'first'
    second = tmp_path / 'second'

    async def run():
        servers = [create_server(Settings(tmp_path, path)) for path in (first, second)]
        async with Client(servers[0], mode=mode) as a, Client(servers[1], mode=mode) as b:
            for client, title in ((a, 'first document'), (b, 'second document')):
                response = await client.call_tool('create_mind_map', {'title': title, 'topics_json': '[]'})
                result = json.loads(response.content[0].text)
                assert result['status'] == 'success'
                assert 'output_path' not in result['data']
            for client, expected in ((a, 'first_document.xmind'), (b, 'second_document.xmind')):
                response = await client.call_tool('list_xmind_files', {})
                result = json.loads(response.content[0].text)
                assert [file['name'] for file in result['data']['files']] == [expected]
    asyncio.run(run())
    assert list(first.glob('*.xmind')) == [first / 'first_document.xmind']
    assert list(second.glob('*.xmind')) == [second / 'second_document.xmind']


def test_alias_normalization_and_unknown_fields_round_trip(tmp_path):
    topics = normalize_topics({'text': 'root', 'customField': {'retain': True}, 'style': {'color': 'red'},
                               'children': {'attached': [{'name': 'normal', 'markers': ['star']}],
                                            'detached': [{'title': 'floating'}]}})
    document = Document(topics[0], extra={'theme': {'custom': True}, 'unknownSheetField': ['keep']})
    path = tmp_path / 'roundtrip.xmind'
    create_xmind_file(document, path)
    loaded = XMindReader(path).read().document
    exported = document_to_sheet(loaded)
    assert exported['theme'] == {'custom': True}
    assert exported['unknownSheetField'] == ['keep']
    root = exported['rootTopic']
    assert root['customField'] == {'retain': True}
    assert root['style'] == {'color': 'red'}
    assert root['children']['attached'][0]['markers'] == ['star']
    assert root['children']['detached'][0]['title'] == 'floating'


@pytest.mark.parametrize('mode', ['auto', 'legacy'])
def test_translation_provider_is_injected_and_archive_resources_are_preserved(tmp_path, mode):
    import zipfile
    source = tmp_path / 'source.xmind'
    create_xmind_file(Document(Topic('中文标题')), source)
    with zipfile.ZipFile(source, 'a') as archive:
        archive.writestr('attachments/original.bin', b'original resource')
        archive.writestr('custom.data', b'unknown field')
    languages = []

    class FakeTranslator:
        def translate(self, text):
            return {'中文标题': 'Translated title'}.get(text, text)

    def factory(language):
        languages.append(language)
        return FakeTranslator()

    service = XMindService(Settings(tmp_path, tmp_path / 'translated'), translator_factory=factory)
    result = service.translate_xmind_titles(str(source), target_lang='en')
    assert result['status'] == 'success'
    assert languages == ['en']
    with zipfile.ZipFile(result['data']['output_file']) as archive:
        assert archive.read('attachments/original.bin') == b'original resource'
        assert archive.read('custom.data') == b'unknown field'
        assert json.loads(archive.read('content.json'))[0]['rootTopic']['title'] == 'Translated title'

    from mcp import Client
    from xmind_mcp.server import create_server
    output = tmp_path / 'mcp-translated.xmind'

    async def run():
        async with Client(create_server(service=service), mode=mode) as client:
            response = await client.call_tool('translate_xmind_titles', {
                'source_filepath': str(source), 'output_filepath': str(output), 'target_lang': 'en'})
            translated = json.loads(response.content[0].text)
            assert translated['status'] == 'success'
            assert translated['data'] == {'filename': output.name}
    asyncio.run(run())
    assert languages == ['en', 'en']
    with zipfile.ZipFile(output) as archive:
        assert archive.read('attachments/original.bin') == b'original resource'
        assert json.loads(archive.read('content.json'))[0]['rootTopic']['title'] == 'Translated title'


def test_cli_and_service_share_conversion_and_explicit_output(tmp_path):
    source = tmp_path / 'input.md'
    source.write_text('# Root\n## Child\n')
    result = subprocess.run([sys.executable, '-c',
                             'from xmind_mcp.cli import convert_main; raise SystemExit(convert_main())',
                             'input.md', 'nested/cli.xmind'], cwd=tmp_path,
                            capture_output=True, text=True, timeout=20)
    assert result.returncode == 0, result.stderr
    service = XMindService(Settings(tmp_path, tmp_path / 'service'))
    assert service.convert_to_xmind(str(source))['status'] == 'success'
    cli_root = service.read_xmind_file(str(tmp_path / 'nested/cli.xmind'))['data']['structure']
    service_root = service.read_xmind_file(str(tmp_path / 'service/input.xmind'))['data']['structure']
    assert cli_root['title'] == service_root['title'] == 'Root'
    assert cli_root['children'][0]['title'] == service_root['children'][0]['title'] == 'Child'


def test_invalid_output_is_reported_without_creating_a_file(tmp_path):
    from xmind_mcp.server import tool_handlers
    service = XMindService(Settings(tmp_path))
    handler = next(fn for fn in tool_handlers(service) if fn.__name__ == 'create_mind_map')
    result = json.loads(handler('title', '[]', 'relative.xmind'))
    assert result['status'] == 'error'
    assert result['error_code'] == 'invalid_input'
    assert not (tmp_path / 'relative.xmind').exists()


@pytest.mark.parametrize('extension', ['txt', 'docx', 'xlsx'])
def test_heading_adapters_share_branch_stack(tmp_path, extension):
    source = tmp_path / f'branches.{extension}'
    headings = [(1, 'A'), (2, 'A1'), (1, 'B'), (3, 'B child')]
    if extension == 'txt':
        source.write_text('Root\n' + ''.join('    ' * level + title + '\n' for level, title in headings))
    elif extension == 'docx':
        from docx import Document as WordDocument
        doc = WordDocument()
        for level, title in headings:
            doc.add_heading(title, level=level)
        doc.save(source)
    else:
        from openpyxl import Workbook
        workbook = Workbook()
        for level, title in headings:
            workbook.active.append([level, title])
        workbook.save(source)
        workbook.close()
    document = ParserRegistry().get_parser(source).parse()
    a, b = document.root.children
    assert a.title == 'A' and b.title == 'B'
    assert a.children[0].title == 'A1'
    assert a.children[0].children == []
    assert b.children[0].title == 'B child'


def test_writer_preserves_default_xmind_presentation(tmp_path):
    from xmind_mcp.parsers.markdown import MarkdownParser
    source = tmp_path / 'presentation.md'
    source.write_text('# Root\n## Child\n')
    sheet = document_to_sheet(MarkdownParser(source).parse())
    assert sheet['class'] == 'sheet'
    assert sheet['rootTopic']['class'] == 'topic'
    assert sheet['rootTopic']['structureClass'] == 'org.xmind.ui.logic.right'
